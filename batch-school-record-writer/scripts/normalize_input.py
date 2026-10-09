#!/usr/bin/env python3
"""Normalize student activity tables and text into a stable JSON structure."""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from pathlib import Path
from typing import Any


ALIASES = {
    "student_id": ["student_id", "학번", "번호", "학생id", "학생번호"],
    "name": ["name", "이름", "성명", "학생명"],
    "grade": ["grade", "학년"],
    "area": ["area", "영역", "기록영역", "학생부영역"],
    "subject": ["subject", "과목", "교과", "교과명"],
    "activity": ["activity", "활동", "활동내용", "주제", "탐구주제"],
    "role": ["role", "역할", "담당", "맡은일"],
    "process": ["process", "과정", "수행과정", "활동과정"],
    "outcome": ["outcome", "결과", "결과물", "산출물"],
    "teacher_observation": [
        "teacher_observation",
        "교사관찰",
        "관찰",
        "관찰내용",
        "교사관찰내용",
    ],
    "student_statement": [
        "student_statement",
        "학생진술",
        "학생소감",
        "자기평가",
        "소감",
    ],
    "target_characters": [
        "target_characters",
        "목표글자수",
        "목표분량",
        "글자수",
    ],
    "existing_text": ["existing_text", "기존문장", "원문", "검토문장"],
}

EVIDENCE_FIELDS = (
    "activity",
    "role",
    "process",
    "outcome",
    "teacher_observation",
    "student_statement",
)

EVIDENCE_LABELS = {
    "activity": "활동",
    "role": "역할",
    "process": "과정",
    "outcome": "결과",
    "teacher_observation": "교사관찰",
    "student_statement": "학생진술",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    if text.endswith(".0") and re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def header_key(value: Any) -> str:
    return re.sub(r"[\s_\-·/()]+", "", clean(value).lower())


ALIAS_LOOKUP = {
    header_key(alias): canonical
    for canonical, aliases in ALIASES.items()
    for alias in aliases
}


def read_text_with_fallback(path: Path) -> str:
    errors = []
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc}")
    raise ValueError("텍스트 인코딩을 읽을 수 없습니다. " + " | ".join(errors))


def read_delimited(path: Path) -> list[dict[str, str]]:
    text = read_text_with_fallback(path)
    sample = text[:4096]
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;").delimiter
    except csv.Error:
        pass
    return [
        {clean(key): clean(value) for key, value in row.items() if key is not None}
        for row in csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
        if any(clean(value) for value in row.values())
    ]


def read_xlsx(path: Path, sheet_name: str | None) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "XLSX 입력에는 openpyxl이 필요합니다. CSV로 저장하거나 openpyxl을 설치하세요."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"시트 '{sheet_name}'을 찾을 수 없습니다: {', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[sheet_name]
    else:
        visible = [
            workbook[name]
            for name in workbook.sheetnames
            if workbook[name].sheet_state == "visible"
        ]
        if not visible:
            raise ValueError("표시 상태인 시트가 없습니다.")
        sheet = visible[0]

    values = sheet.iter_rows(values_only=True)
    try:
        headers = [clean(value) for value in next(values)]
    except StopIteration:
        return []

    rows = []
    for cells in values:
        row = {
            header: clean(value)
            for header, value in zip(headers, cells)
            if header
        }
        if any(row.values()):
            rows.append(row)
    return rows


def read_json(path: Path) -> list[dict[str, Any]]:
    data = json.loads(read_text_with_fallback(path))
    if isinstance(data, dict):
        data = data.get("records", [data])
    if not isinstance(data, list) or not all(isinstance(item, dict) for item in data):
        raise ValueError("JSON은 객체 배열 또는 records 배열이어야 합니다.")
    return data


def parse_markdown_table(text: str) -> list[dict[str, str]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for index in range(len(lines) - 1):
        if "|" not in lines[index]:
            continue
        separator = lines[index + 1].strip().strip("|")
        if not re.fullmatch(r"\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+", separator):
            continue
        headers = [cell.strip() for cell in lines[index].strip("|").split("|")]
        rows = []
        for line in lines[index + 2 :]:
            if "|" not in line:
                break
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            row = {
                header: cells[pos] if pos < len(cells) else ""
                for pos, header in enumerate(headers)
                if header
            }
            if any(row.values()):
                rows.append(row)
        return rows
    return []


def parse_key_value_text(text: str) -> list[dict[str, str]]:
    chunks = re.split(
        r"\n\s*---+\s*\n|\n\s*\n(?=\s*(?:학번|student_id|이름|name)\s*[:：])",
        text.strip(),
        flags=re.IGNORECASE,
    )
    rows = []
    for chunk in chunks:
        row: dict[str, str] = {}
        current_key = ""
        for raw_line in chunk.splitlines():
            line = raw_line.strip()
            if not line:
                continue
            match = re.match(r"^([^:：]+)[:：]\s*(.*)$", line)
            if match:
                current_key = clean(match.group(1))
                row[current_key] = clean(match.group(2))
            elif current_key:
                row[current_key] = clean(row[current_key] + "\n" + line)
            else:
                row.setdefault("활동내용", "")
                row["활동내용"] = clean(row["활동내용"] + "\n" + line)
        if any(row.values()):
            rows.append(row)
    return rows


def read_text_records(path: Path) -> list[dict[str, str]]:
    text = read_text_with_fallback(path)
    table = parse_markdown_table(text)
    return table if table else parse_key_value_text(text)


def load_rows(path: Path, sheet_name: str | None) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        return read_delimited(path)
    if suffix == ".xlsx":
        return read_xlsx(path, sheet_name)
    if suffix == ".json":
        return read_json(path)
    if suffix in {".txt", ".md", ".markdown"}:
        return read_text_records(path)
    raise ValueError("지원 형식: .xlsx, .csv, .tsv, .json, .txt, .md")


def canonicalize_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    canonical = {field: "" for field in ALIASES}
    source_fields: dict[str, str] = {}

    for raw_key, raw_value in row.items():
        key = clean(raw_key)
        value = clean(raw_value)
        if not key:
            continue
        mapped = ALIAS_LOOKUP.get(header_key(key))
        if mapped and not canonical[mapped]:
            canonical[mapped] = value
        else:
            source_fields[key] = value

    evidence_parts = []
    for field in EVIDENCE_FIELDS:
        if canonical[field]:
            evidence_parts.append(f"{EVIDENCE_LABELS[field]}: {canonical[field]}")

    record_key = canonical["student_id"] or canonical["name"] or f"row-{row_number}"
    return {
        "record_key": record_key,
        **canonical,
        "evidence_text": " | ".join(evidence_parts),
        "source_fields": source_fields,
        "source_row": row_number,
    }


def build_warnings(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    warnings: list[dict[str, Any]] = []
    seen: dict[str, int] = {}
    for record in records:
        row = record["source_row"]
        student_id = record["student_id"]
        if student_id:
            if student_id in seen:
                warnings.append(
                    {
                        "row": row,
                        "code": "duplicate_student_id",
                        "message": f"학번 {student_id}이(가) {seen[student_id]}행과 중복됩니다.",
                    }
                )
            else:
                seen[student_id] = row
        if not record["student_id"] and not record["name"]:
            warnings.append(
                {
                    "row": row,
                    "code": "missing_identity",
                    "message": "학번과 이름이 모두 없습니다.",
                }
            )
        if not record["area"]:
            warnings.append(
                {
                    "row": row,
                    "code": "missing_area",
                    "message": "기록 영역이 없습니다.",
                }
            )
        if not record["evidence_text"] and not record["existing_text"]:
            warnings.append(
                {
                    "row": row,
                    "code": "missing_evidence",
                    "message": "활동·역할·과정·결과·관찰 근거가 없습니다.",
                }
            )
    return warnings


def main() -> int:
    parser = argparse.ArgumentParser(
        description="학생별 표·엑셀·CSV·텍스트를 표준 JSON으로 정규화합니다."
    )
    parser.add_argument("input", type=Path, help="입력 파일")
    parser.add_argument("--sheet", help="XLSX 시트 이름")
    parser.add_argument("--output", type=Path, help="출력 JSON 파일. 생략하면 표준 출력")
    args = parser.parse_args()

    try:
        rows = load_rows(args.input, args.sheet)
        records = [
            canonicalize_row(row, index + 2)
            for index, row in enumerate(rows)
        ]
        result = {
            "source": str(args.input),
            "count": len(records),
            "warnings": build_warnings(records),
            "records": records,
        }
        payload = json.dumps(result, ensure_ascii=False, indent=2)
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(payload + "\n", encoding="utf-8")
        else:
            sys.stdout.write(payload + "\n")
        return 0
    except Exception as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
